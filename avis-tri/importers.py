"""Lecteurs de fichiers ; aucune exécution de formule ou de macro."""
import csv
import io
import json
import re
import zipfile
from datetime import date, datetime
from itertools import islice
from pathlib import Path

MAX_ROWS = 500_000
MAX_COLUMNS = 500


def cell(value):
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def table(rows, header=True):
    rows = iter(rows)
    first = next(rows, [])
    if not header:
        import itertools
        rows = itertools.chain([first], rows)
        first = [f'colonne_{i+1}' for i in range(len(first))]
    columns, seen = [], set()
    for i, value in enumerate(first):
        base = cell(value).strip() or f'colonne_{i+1}'
        name, n = base, 2
        while name in seen:
            name, n = f'{base}_{n}', n + 1
        seen.add(name)
        columns.append(name)
    if len(columns) > MAX_COLUMNS:
        raise ValueError(f'Maximum : {MAX_COLUMNS} colonnes.')
    records = []
    for index, row in enumerate(rows, 2 if header else 1):
        values = [cell(v) for v in row]
        if not any(v.strip() for v in values):
            continue
        if len(values) > len(columns):
            if any(values[len(columns):]):
                raise ValueError(f'Ligne {index} : trop de colonnes. Choisis le séparateur dans les options d’import.')
            values = values[:len(columns)]
        records.append(dict(zip(columns, values + [''] * (len(columns) - len(values)))))
        if len(records) > MAX_ROWS:
            raise ValueError(f'Maximum : {MAX_ROWS:,} avis ; scinde le fichier.')
    return columns, records


def zip_guard(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if sum(x.file_size for x in z.infolist()) > 300 * 1024**2:
            raise ValueError('Archive décompressée supérieure à 300 Mo.')


def decode(data, encoding='auto'):
    if encoding != 'auto':
        if encoding not in {'utf-8-sig', 'utf-16', 'cp1252', 'latin-1'}:
            raise ValueError('Encodage non pris en charge.')
        return data.decode(encoding)
    if data[:2] in (b'\xff\xfe', b'\xfe\xff'):
        return data.decode('utf-16')
    try:
        return data.decode('utf-8-sig')
    except UnicodeDecodeError:
        return data.decode('cp1252')


def read_file(data, filename, options=None):
    options = options or {}
    ext = Path(filename).suffix.lower()
    info, sheets = '', []
    header = options.get('header', True)
    if ext in {'.xlsx', '.xlsm', '.xls'}:
        if ext == '.xls':
            import xlrd
            book = xlrd.open_workbook(file_contents=data, on_demand=True)
            sheets = book.sheet_names()
            name = options.get('sheet') or sheets[0]
            sheet = book.sheet_by_name(name)
            def xls_rows():
                for i in range(sheet.nrows):
                    yield [xlrd.xldate_as_datetime(c.value, book.datemode) if c.ctype == xlrd.XL_CELL_DATE else c.value for c in sheet.row(i)]
            try:
                columns, records = table(xls_rows(), header)
            finally:
                book.release_resources()
        else:
            zip_guard(data)
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            try:
                sheets = book.sheetnames
                name = options.get('sheet') or sheets[0]
                if name not in sheets:
                    raise ValueError('Feuille inconnue.')
                sheet = book[name]
                # Les dimensions d’un classeur peuvent être incorrectes ; le lecteur suit les cellules réelles.
                sheet.reset_dimensions()
                columns, records = table(sheet.iter_rows(values_only=True), header)
            finally:
                book.close()
        info = f'Feuille : {name}. Dates converties en ISO. Formules : résultats enregistrés dans le classeur uniquement ; aucune formule calculée.'
    elif ext == '.docx':
        zip_guard(data)
        from docx import Document
        doc = Document(io.BytesIO(data))
        rows = [[p.text] for p in doc.paragraphs if p.text.strip()]
        for t in doc.tables:
            for r in t.rows:
                rows.append([' | '.join(c.text for c in r.cells)])
        columns, records = table(iter([['texte']] + rows))
        info = 'Word : un paragraphe ou une ligne de tableau par avis. Vérifie la segmentation dans l’aperçu.'
    elif ext == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError('PDF protégé : fournis une copie déverrouillée.')
        records, missing = [], []
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ''
            if not text.strip():
                missing.append(i)
            for part in re.split(r'\n\s*\n', text):
                if part.strip():
                    records.append({'texte': part.strip(), 'page': str(i)})
            if len(records) > MAX_ROWS:
                raise ValueError('Trop de paragraphes dans le PDF.')
        if not records:
            raise ValueError('PDF sans texte extractible (scan probable). Effectue un OCR avant l’import.')
        columns = ['texte', 'page']
        info = 'PDF : segmentation par paragraphes à vérifier, les mises en page complexes peuvent mélanger plusieurs avis.'
        if missing:
            info += f' Pages sans texte extractible : {", ".join(map(str, missing[:30]))}. OCR nécessaire pour ces pages.'
    else:
        text = decode(data, options.get('encoding', 'auto'))
        if ext in {'.json', '.jsonl', '.ndjson'}:
            obj = [json.loads(line) for line in text.splitlines() if line.strip()] if ext != '.json' else json.loads(text)
            if isinstance(obj, dict):
                candidates = [obj[k] for k in ['reviews', 'avis', 'data', 'records'] if isinstance(obj.get(k), list)]
                if len(candidates) != 1:
                    raise ValueError('JSON : choisis une liste d’objets, ou un objet contenant une seule liste reviews/avis/data/records.')
                obj = candidates[0]
            if not isinstance(obj, list) or not all(isinstance(r, dict) for r in obj):
                raise ValueError('JSON attendu : une liste d’objets, un objet par avis.')
            if len(obj) > MAX_ROWS:
                raise ValueError('Plus de 500 000 avis.')
            columns = list(dict.fromkeys(k for r in obj for k in r))
            if len(columns) > MAX_COLUMNS:
                raise ValueError('Plus de 500 colonnes.')
            records = [{k: cell(r.get(k)) for k in columns} for r in obj]
            info = 'JSON importé ; les objets imbriqués restent du JSON dans leur colonne source.'
        elif ext == '.txt' and options.get('text_mode', 'lines') != 'table':
            parts = re.split(r'\n\s*\n', text) if options.get('text_mode') == 'paragraphs' else text.splitlines()
            columns, records = table(iter([['texte']] + [[p.strip()] for p in parts if p.strip()]))
            info = 'Texte libre : un avis par ' + ('paragraphe.' if options.get('text_mode') == 'paragraphs' else 'ligne.')
        elif ext in {'.csv', '.tsv', '.txt'}:
            csv.field_size_limit(5 * 1024**2)
            delimiter = options.get('delimiter', 'auto')
            if delimiter == 'auto':
                try:
                    delimiter = csv.Sniffer().sniff(text[:65536], delimiters=',;\t|').delimiter
                except csv.Error:
                    delimiter = '\t' if ext == '.tsv' else ','
            if delimiter not in {',', ';', '\t', '|'}:
                raise ValueError('Séparateur non pris en charge.')
            columns, records = table(csv.reader(io.StringIO(text, newline=''), delimiter=delimiter, doublequote=True), header)
            info = f'Séparateur : {repr(delimiter)}. Encodage : {options.get("encoding", "auto")}.'
        else:
            raise ValueError('Formats : CSV, TSV, TXT, XLSX, XLSM, XLS, JSON, JSONL, DOCX, PDF texte.')
    if not records and not sheets:
        raise ValueError('Aucun avis dans cette feuille ou ce fichier. Vérifie les options d’import.')
    if not records and sheets:
        info += ' Feuille vide : choisis une autre feuille dans le sélecteur.'
    return {'columns': columns, 'records': records, 'info': info, 'sheets': sheets}
