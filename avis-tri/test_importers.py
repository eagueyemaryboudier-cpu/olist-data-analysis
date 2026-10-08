import io
import json
import unittest
from datetime import datetime
import server  # active les bibliothèques locales
from importers import read_file


class ImportTests(unittest.TestCase):
    def test_text_modes(self):
        data = b'Premier avis\nSuite\n\nSecond avis'
        self.assertEqual(len(read_file(data, 'a.txt')['records']), 3)
        self.assertEqual(len(read_file(data, 'a.txt', {'text_mode': 'paragraphs'})['records']), 2)

    def test_csv_override_and_no_header(self):
        result = read_file('café|5\r\néchec|1'.encode('utf-16'), 'a.csv', {'delimiter': '|', 'header': False})
        self.assertEqual(result['columns'], ['colonne_1', 'colonne_2'])
        self.assertEqual(result['records'][0]['colonne_1'], 'café')

    def test_jsonl_nested(self):
        result = read_file(b'{"text":"ok", "meta":{"a":1}}\n{"text":"bad"}', 'a.jsonl')
        self.assertEqual(json.loads(result['records'][0]['meta']), {'a': 1})
        wrapped = read_file(b'{"reviews":[{"text":"ok"}]}', 'a.json')
        self.assertEqual(len(wrapped['records']), 1)

    def test_more_than_old_row_limit(self):
        result = read_file(b'text\n' + b'ok\n' * 100_001, 'a.csv')
        self.assertEqual(len(result['records']), 100_001)

    def test_workbook_select_sheet_dates_and_formula(self):
        from openpyxl import Workbook
        book = Workbook()
        book.active.append(['ignore'])
        sheet = book.create_sheet('Avis')
        sheet.append(['commentaire', 'date', 'note'])
        sheet.append(['Très bien', datetime(2024, 1, 2), '=1+4'])
        buffer = io.BytesIO()
        book.save(buffer)
        result = read_file(buffer.getvalue(), 'a.xlsx', {'sheet': 'Avis'})
        self.assertEqual(result['sheets'], ['Sheet', 'Avis'])
        self.assertEqual(result['records'][0]['date'], '2024-01-02T00:00:00')
        self.assertEqual(result['records'][0]['note'], '')  # pas de calcul caché

    def test_empty_first_sheet_can_select_another(self):
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.create_sheet('Avis')
        sheet.append(['texte'])
        sheet.append(['Excellent'])
        b = io.BytesIO()
        book.save(b)
        result = read_file(b.getvalue(), 'a.xlsx')
        self.assertEqual(result['records'], [])
        self.assertIn('Avis', result['sheets'])

    def test_docx_paragraph_and_table(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph('Très bon produit')
        table = doc.add_table(rows=1, cols=2)
        table.cell(0,0).text = 'Livraison en retard'
        table.cell(0,1).text = '1'
        b = io.BytesIO()
        doc.save(b)
        result = read_file(b.getvalue(), 'a.docx')
        self.assertEqual(len(result['records']), 2)
        self.assertIn('Livraison en retard', result['records'][1]['texte'])

    def test_pdf_no_silent_empty_scan(self):
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        b = io.BytesIO()
        writer.write(b)
        with self.assertRaisesRegex(ValueError, 'OCR'):
            read_file(b.getvalue(), 'scan.pdf')

    def test_non_latin_is_not_empty(self):
        self.assertEqual(server.classify('商品很好')[0], 'Indéterminé')

    def test_custom_theme_expression(self):
        result = server.analyze([{'text':'Le service après vente est excellent'}], {'text':'text','themes':{'Assistance':['service après vente']}})
        self.assertEqual(result[0]['themes'], ['Assistance'])


if __name__ == '__main__':
    unittest.main()
